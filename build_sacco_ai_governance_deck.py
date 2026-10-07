"""
Builds an Executive PowerPoint Presentation (.pptx) for SACCO Boards:
"AI Governance & Responsible Innovation Framework for SACCOs"
Now updated with 13 Comprehensive Slides, including Mobile App Ecosystem Integration for Banki Kuu SACCO.
Tailored for SASRA & CBK Compliance under KDPA 2019.
"""

import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    COLOR_NAVY_DARK = RGBColor(9, 31, 61)       # #091F3D
    COLOR_NAVY_CARD = RGBColor(15, 42, 79)      # #0F2A4F
    COLOR_GOLD = RGBColor(245, 197, 66)         # #F5C542
    COLOR_GOLD_DARK = RGBColor(212, 140, 20)    # #D48C14
    COLOR_CYAN = RGBColor(0, 242, 254)          # #00F2FE
    COLOR_EMERALD = RGBColor(16, 185, 129)      # #10B981
    COLOR_WHITE = RGBColor(255, 255, 255)
    COLOR_TEXT_MUTED = RGBColor(203, 213, 225)  # #CBD5E1
    COLOR_CARD_BORDER = RGBColor(30, 64, 105)

    def set_slide_background(slide, color):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.color.rgb = color
        return bg

    def add_header(slide, category_text, title_text):
        badge_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.35))
        tf = badge_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = category_text.upper()
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = COLOR_GOLD
        
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.733), Inches(0.6))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(22)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_WHITE

    # --------------------------------------------------------------------------
    # SLIDE 1: TITLE SLIDE
    # --------------------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1, COLOR_NAVY_DARK)

    card1 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.0), Inches(11.333), Inches(5.5))
    card1.fill.solid()
    card1.fill.fore_color.rgb = COLOR_NAVY_CARD
    card1.line.color.rgb = COLOR_GOLD
    card1.line.width = Pt(2)

    tb = slide1.shapes.add_textbox(Inches(1.5), Inches(1.5), Inches(10.333), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "🏛️ BANKI KUU SACCO EXECUTIVE BOARDROOM BRIEFING"
    p0.font.size = Pt(12)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_GOLD

    p1 = tf.add_paragraph()
    p1.text = "AI Governance & Mobile App Integration Strategy for SACCOs"
    p1.font.size = Pt(30)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_WHITE
    p1.space_before = Pt(14)
    p1.space_after = Pt(14)

    p2 = tf.add_paragraph()
    p2.text = "Enhancing Banki Kuu SACCO's Existing Mobile App with AI Underwriting, In-App E-Voting, SASRA & CBK Compliance, and KDPA 2019 Safeguards"
    p2.font.size = Pt(14)
    p2.font.color.rgb = COLOR_TEXT_MUTED
    p2.space_after = Pt(28)

    p3 = tf.add_paragraph()
    p3.text = "Prepared for: Board of Directors & Executive Management | Strategy Series 2026"
    p3.font.size = Pt(12)
    p3.font.bold = True
    p3.font.color.rgb = COLOR_CYAN

    # --------------------------------------------------------------------------
    # SLIDE 2: EXECUTIVE SUMMARY & STRATEGIC RATIONALE
    # --------------------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2, COLOR_NAVY_DARK)
    add_header(slide2, "Executive Context", "Why SACCOs Require Dedicated AI Governance")

    box_l = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    box_l.fill.solid()
    box_l.fill.fore_color.rgb = COLOR_NAVY_CARD
    box_l.line.color.rgb = COLOR_CARD_BORDER

    tb_l = slide2.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(5.0), Inches(4.8))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True

    p = tf_l.paragraphs[0]
    p.text = "🚀 High-Impact AI Opportunities in SACCOs"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_GOLD

    opps = [
        ("Automated Credit Underwriting", "Instant scoring for instant emergency & salary advance loans."),
        ("Predictive Fraud Detection", "Real-time detection of suspicious M-Pesa & banking withdrawals."),
        ("Member Churn & Liquidity AI", "Predictive forecasting of dividend expectations & deposit trends."),
        ("24/7 Member Conversational AI", "Automated WhatsApp & web assistants for loan inquiries.")
    ]
    for title, desc in opps:
        p_t = tf_l.add_paragraph()
        p_t.text = f"• {title}: "
        p_t.font.bold = True
        p_t.font.size = Pt(12)
        p_t.font.color.rgb = COLOR_WHITE
        p_t.space_before = Pt(8)
        
        p_d = tf_l.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = COLOR_TEXT_MUTED

    box_r = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    box_r.fill.solid()
    box_r.fill.fore_color.rgb = COLOR_NAVY_CARD
    box_r.line.color.rgb = COLOR_GOLD

    tb_r = slide2.shapes.add_textbox(Inches(7.1), Inches(1.8), Inches(5.1), Inches(4.8))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True

    p = tf_r.paragraphs[0]
    p.text = "⚠️ Governance Mandatory Imperatives"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_GOLD

    risks = [
        ("Algorithmic Bias Risk", "Preventing unfair loan rejections based on proxy data demographics."),
        ("SASRA Regulatory Compliance", "Ensuring loan decisions can be audited by regulatory scrutinizers."),
        ("Data Sovereignty & KDPA 2019", "Strictly protecting member PII from unencrypted AI cloud leaks."),
        ("Reputational & Fiduciary Duty", "Safeguarding cooperative member trust and financial equity.")
    ]
    for title, desc in risks:
        p_t = tf_r.add_paragraph()
        p_t.text = f"• {title}: "
        p_t.font.bold = True
        p_t.font.size = Pt(12)
        p_t.font.color.rgb = COLOR_CYAN
        p_t.space_before = Pt(8)

        p_d = tf_r.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = COLOR_TEXT_MUTED

    # --------------------------------------------------------------------------
    # SLIDE 3: TOP 5 HIGH-IMPACT AI QUICK WINS FOR BANKI KUU SACCO
    # --------------------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3, COLOR_NAVY_DARK)
    add_header(slide3, "Strategic AI Quick Wins", "Top 5 High-Impact AI Quick Wins for Banki Kuu SACCO")

    quick_wins_data = [
        ("1. 📱 24/7 WhatsApp & In-App Assistant", "30 Days", "High Impact", "70% Reduction in Call Center Support Load", COLOR_GOLD),
        ("2. 🎟️ In-App E-Voting & STRIDE™ AGM", "Immediate", "High Impact", "95%+ Savings on AGM & Ballot Expenses", COLOR_CYAN),
        ("3. ⚡ AI Underwriting for App Loans", "45 Days", "High Impact", "+30% Growth in Micro-Loan Interest Income", COLOR_EMERALD),
        ("4. 🛡️ Real-Time App Fraud Anomaly AI", "30 Days", "High Impact", "Prevents M-Pesa Fraud & Account Takeovers", COLOR_GOLD),
        ("5. 📄 Intelligent Document OCR", "60 Days", "Medium Impact", "Reduces Guarantor Approval from 3 Days to 5 Mins", COLOR_CYAN)
    ]

    for i, (qw_t, qw_speed, qw_imp, qw_roi, qw_color) in enumerate(quick_wins_data):
        y_pos = Inches(1.6 + i * 1.05)
        box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y_pos, Inches(11.733), Inches(0.9))
        box.fill.solid()
        box.fill.fore_color.rgb = COLOR_NAVY_CARD
        box.line.color.rgb = qw_color
        box.line.width = Pt(1.5)

        tb_qw = slide3.shapes.add_textbox(Inches(1.0), y_pos + Inches(0.1), Inches(11.3), Inches(0.7))
        tf_qw = tb_qw.text_frame
        tf_qw.word_wrap = True

        p = tf_qw.paragraphs[0]
        p.text = f"{qw_t}   |   Speed: {qw_speed}   |   {qw_imp}"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = qw_color

        p_sub = tf_qw.add_paragraph()
        p_sub.text = f"🎯 Measurable ROI & Value: {qw_roi}"
        p_sub.font.size = Pt(11)
        p_sub.font.color.rgb = COLOR_TEXT_MUTED
        p_sub.space_before = Pt(2)

    # --------------------------------------------------------------------------
    # SLIDE 4: QUICK WINS DEEP DIVE — IMMEDIATE OPERATIONAL ROI
    # --------------------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4, COLOR_NAVY_DARK)
    add_header(slide4, "Operational ROI", "Quick Wins Deep-Dive: Immediate Member Value")

    cards_qw = [
        ("📱 24/7 WhatsApp & App Assistant", "• Resolves member queries instantly 24/7 (balances, loan eligibility, dividend status).\n• Verifies identity via SMS OTP & reduces Secretariat calls by 70%.", COLOR_GOLD),
        ("🎟️ STRIDE™ In-App E-Voting & AGM", "• 1-Click QR Accreditation & 2-Second Gate Scan.\n• Encrypted E-Voting with SHA-256 Receipts & NLP Member Pulse Sentiment Analysis.", COLOR_CYAN),
        ("⚡ Instant App Credit Scoring", "• Evaluates deposit history & salary consistency in 3 seconds.\n• Disburses instant emergency loans via M-Pesa with near-zero NPLs.", COLOR_EMERALD)
    ]

    for i, (c_t, c_d, c_col) in enumerate(cards_qw):
        x_pos = Inches(0.8 + i * 4.0)
        box = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x_pos, Inches(1.8), Inches(3.7), Inches(5.0))
        box.fill.solid()
        box.fill.fore_color.rgb = COLOR_NAVY_CARD
        box.line.color.rgb = c_col
        box.line.width = Pt(1.5)

        tb_qwc = slide4.shapes.add_textbox(x_pos + Inches(0.2), Inches(2.0), Inches(3.3), Inches(4.5))
        tf_qwc = tb_qwc.text_frame
        tf_qwc.word_wrap = True

        p = tf_qwc.paragraphs[0]
        p.text = c_t
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = c_col

        p_d = tf_qwc.add_paragraph()
        p_d.text = c_d
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = COLOR_TEXT_MUTED
        p_d.space_before = Pt(12)

    # --------------------------------------------------------------------------
    # SLIDE 5 (NEW): MOBILE APP ECOSYSTEM INTEGRATION
    # --------------------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5, COLOR_NAVY_DARK)
    add_header(slide5, "Ecosystem Integration", "Enhancing Banki Kuu SACCO's Existing Mobile App")

    app_modules = [
        ("📱 In-App E-Voting Tile (Single Sign-On)", "Members open their existing Banki Kuu SACCO phone app and tap '🎟️ 58th AGM & E-Voting' to unlock their ballot without re-entering credentials.", COLOR_GOLD),
        ("⚡ AI Loan Underwriting Engine API", "Enhances the existing 'Borrow Loan' app feature with 3-second machine learning credit scoring and instant M-Pesa disbursement.", COLOR_CYAN),
        ("🤖 In-App AI Member Assistant", "Integrates an AI chatbot inside the app for instant loan eligibility checks, dividend statements, and balance inquiries.", COLOR_EMERALD),
        ("🛡️ Real-Time App Anomaly & Fraud Guard", "Monitors device fingerprints and transaction patterns to prevent unauthorized app loan withdrawals.", COLOR_GOLD)
    ]

    for idx, (m_t, m_d, m_c) in enumerate(app_modules):
        row = idx // 2
        col = idx % 2
        x = Inches(0.8 + col * 5.9)
        y = Inches(1.8 + row * 2.5)

        box = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.6), Inches(2.2))
        box.fill.solid()
        box.fill.fore_color.rgb = COLOR_NAVY_CARD
        box.line.color.rgb = m_c

        tb_m = slide5.shapes.add_textbox(x + Inches(0.2), y + Inches(0.2), Inches(5.2), Inches(1.8))
        tf_m = tb_m.text_frame
        tf_m.word_wrap = True

        p = tf_m.paragraphs[0]
        p.text = m_t
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = m_c

        p_d = tf_m.add_paragraph()
        p_d.text = m_d
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = COLOR_TEXT_MUTED
        p_d.space_before = Pt(6)

    # --------------------------------------------------------------------------
    # SLIDE 6: THE 5 PILLARS OF SACCO AI GOVERNANCE
    # --------------------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6, COLOR_NAVY_DARK)
    add_header(slide6, "Strategic Framework", "The 5 Pillars of SACCO AI Governance")

    pillars = [
        ("Pillar 1", "Regulatory Alignment", "SASRA & CBK Risk Guidelines, Model Audits", COLOR_GOLD),
        ("Pillar 2", "Data Protection", "KDPA 2019, PII Encryption, Explicit Consent", COLOR_CYAN),
        ("Pillar 3", "Algorithmic Fairness", "Bias Mitigation, Equal Odds Underwriting", COLOR_EMERALD),
        ("Pillar 4", "Explainable AI (XAI)", "Member Right to Explanation, SHAP Analysis", COLOR_GOLD),
        ("Pillar 5", "Continuous MLOps", "Model Drift Monitoring, Cyber Resilience", COLOR_CYAN)
    ]

    col_width = Inches(2.2)
    gap = Inches(0.2)
    start_x = Inches(0.8)

    for idx, (p_num, p_name, p_desc, p_col) in enumerate(pillars):
        x_pos = start_x + idx * (col_width + gap)
        box = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x_pos, Inches(1.8), col_width, Inches(5.0))
        box.fill.solid()
        box.fill.fore_color.rgb = COLOR_NAVY_CARD
        box.line.color.rgb = p_col
        box.line.width = Pt(1.5)

        tb_p = slide6.shapes.add_textbox(x_pos + Inches(0.15), Inches(2.0), col_width - Inches(0.3), Inches(4.5))
        tf_p = tb_p.text_frame
        tf_p.word_wrap = True

        p = tf_p.paragraphs[0]
        p.text = p_num
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = p_col

        p_t = tf_p.add_paragraph()
        p_t.text = p_name
        p_t.font.size = Pt(15)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_WHITE
        p_t.space_before = Pt(6)
        p_t.space_after = Pt(12)

        p_d = tf_p.add_paragraph()
        p_d.text = p_desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = COLOR_TEXT_MUTED

    # --------------------------------------------------------------------------
    # SLIDE 7: PILLAR 1 — REGULATORY ALIGNMENT (SASRA & CBK)
    # --------------------------------------------------------------------------
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7, COLOR_NAVY_DARK)
    add_header(slide7, "Pillar 1", "Regulatory Alignment: SASRA & CBK Compliance")

    cards_data = [
        ("📋 Model Risk Inventory", "Catalog all algorithmic models deployed across SACCO operations into High, Medium, and Low risk tiers.", COLOR_GOLD),
        ("🛡️ SASRA Audit Readiness", "Ensure all automated credit decision trees produce deterministic, reproducible audit logs for regulators.", COLOR_CYAN),
        ("👨‍💼 Independent Validation", "Conduct annual third-party model risk audits before deploying AI models for member financial services.", COLOR_EMERALD)
    ]

    for i, (c_title, c_desc, c_color) in enumerate(cards_data):
        y_pos = Inches(1.8 + i * 1.7)
        box = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y_pos, Inches(11.733), Inches(1.4))
        box.fill.solid()
        box.fill.fore_color.rgb = COLOR_NAVY_CARD
        box.line.color.rgb = c_color

        tb_c = slide7.shapes.add_textbox(Inches(1.1), y_pos + Inches(0.2), Inches(11.1), Inches(1.0))
        tf_c = tb_c.text_frame
        tf_c.word_wrap = True

        p = tf_c.paragraphs[0]
        p.text = c_title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = c_color

        p_d = tf_c.add_paragraph()
        p_d.text = c_desc
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = COLOR_TEXT_MUTED
        p_d.space_before = Pt(4)

    # --------------------------------------------------------------------------
    # SLIDE 8: PILLAR 2 — DATA SOVEREIGNTY & PII PROTECTION (KDPA 2019)
    # --------------------------------------------------------------------------
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8, COLOR_NAVY_DARK)
    add_header(slide8, "Pillar 2", "Data Sovereignty & Member PII Protection (KDPA 2019)")

    col1_box = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0))
    col1_box.fill.solid()
    col1_box.fill.fore_color.rgb = COLOR_NAVY_CARD
    col1_box.line.color.rgb = COLOR_CYAN

    tb_c1 = slide8.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.0), Inches(4.5))
    tf_c1 = tb_c1.text_frame
    tf_c1.word_wrap = True

    p = tf_c1.paragraphs[0]
    p.text = "🔒 Legal & Technical Controls"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_CYAN

    controls = [
        ("Kenya Data Protection Act 2019 Compliance", "Strict adherence to ODPC guidelines regarding financial data processing."),
        ("Zero-Trust Member PII Masking", "Automatically mask member National IDs, phone numbers, and names prior to model training."),
        ("On-Premise / Sovereign Hosting", "Keep member telemetry within accredited Kenyan data centers.")
    ]
    for t, d in controls:
        p_t = tf_c1.add_paragraph()
        p_t.text = f"• {t}:"
        p_t.font.bold = True
        p_t.font.size = Pt(12)
        p_t.font.color.rgb = COLOR_WHITE
        p_t.space_before = Pt(8)

        p_d = tf_c1.add_paragraph()
        p_d.text = d
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = COLOR_TEXT_MUTED

    col2_box = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0))
    col2_box.fill.solid()
    col2_box.fill.fore_color.rgb = COLOR_NAVY_CARD
    col2_box.line.color.rgb = COLOR_GOLD

    tb_c2 = slide8.shapes.add_textbox(Inches(7.1), Inches(2.0), Inches(5.1), Inches(4.5))
    tf_c2 = tb_c2.text_frame
    tf_c2.word_wrap = True

    p = tf_c2.paragraphs[0]
    p.text = "📄 Explicit Consent & Data Lifecycle"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_GOLD

    consent = [
        ("Granular Member Opt-In", "Provide clear opt-in checkboxes for automated AI credit profiling."),
        ("Purpose Limitation Principle", "Ensure data collected for loan appraisal is never re-used for external profiling."),
        ("Right to Erasure & Rectification", "Enable members to request updates or deletion of outdated algorithmic inputs.")
    ]
    for t, d in consent:
        p_t = tf_c2.add_paragraph()
        p_t.text = f"• {t}:"
        p_t.font.bold = True
        p_t.font.size = Pt(12)
        p_t.font.color.rgb = COLOR_WHITE
        p_t.space_before = Pt(8)

        p_d = tf_c2.add_paragraph()
        p_d.text = d
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = COLOR_TEXT_MUTED

    # --------------------------------------------------------------------------
    # SLIDE 9: PILLAR 3 — ALGORITHMIC FAIRNESS & ETHICAL UNDERWRITING
    # --------------------------------------------------------------------------
    slide9 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide9, COLOR_NAVY_DARK)
    add_header(slide9, "Pillar 3", "Algorithmic Fairness & Ethical Credit Underwriting")

    quads = [
        ("⚖️ Demographic Parity", "Ensure loan approval rates are statistically balanced across member cohorts without systemic skew.", COLOR_GOLD),
        ("🔍 Disparate Impact Audits", "Measure algorithmic decision ratios continuously to identify and correct unintended bias.", COLOR_CYAN),
        ("🚫 Sensitive Attribute Masking", "Exclude protected categories (gender, marital status, region) from credit scoring models.", COLOR_EMERALD),
        ("🤝 Cooperative Values Alignment", "Preserve SACCO core principles of social equity, member mutual benefit, and financial inclusion.", COLOR_GOLD)
    ]

    for idx, (q_t, q_d, q_c) in enumerate(quads):
        row = idx // 2
        col = idx % 2
        x = Inches(0.8 + col * 5.9)
        y = Inches(1.8 + row * 2.5)

        box = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.6), Inches(2.2))
        box.fill.solid()
        box.fill.fore_color.rgb = COLOR_NAVY_CARD
        box.line.color.rgb = q_c

        tb_q = slide9.shapes.add_textbox(x + Inches(0.2), y + Inches(0.2), Inches(5.2), Inches(1.8))
        tf_q = tb_q.text_frame
        tf_q.word_wrap = True

        p = tf_q.paragraphs[0]
        p.text = q_t
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = q_c

        p_d = tf_q.add_paragraph()
        p_d.text = q_d
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = COLOR_TEXT_MUTED
        p_d.space_before = Pt(6)

    # --------------------------------------------------------------------------
    # SLIDE 10: PILLAR 4 — EXPLAINABLE AI (XAI) & MEMBER RIGHT TO EXPLANATION
    # --------------------------------------------------------------------------
    slide10 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide10, COLOR_NAVY_DARK)
    add_header(slide10, "Pillar 4", "Explainable AI (XAI) & Member Right to Explanation")

    box_main = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.733), Inches(5.0))
    box_main.fill.solid()
    box_main.fill.fore_color.rgb = COLOR_NAVY_CARD
    box_main.line.color.rgb = COLOR_GOLD

    tb_m = slide10.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(11.1), Inches(4.5))
    tf_m = tb_m.text_frame
    tf_m.word_wrap = True

    p = tf_m.paragraphs[0]
    p.text = "💡 Moving from 'Black-Box' Algorithms to Transparent Decision-Making"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_GOLD

    points = [
        ("Member-Facing Credit Explanations", "When a loan is modified or declined, provide plain-language explanations of top contributing factors (e.g. share guarantee ratio, debt-to-income floor)."),
        ("SHAP & LIME Feature Scoring", "Utilize Explainable AI tools (SHAP values) to quantify exact feature weights behind every credit score recommendation."),
        ("Human-in-the-Loop Appeals Process", "Establish a formal SACCO Credit Committee override workflow for members to request human re-appraisal."),
        ("Transparent Interest & Fee Calculations", "Provide clear visibility into how risk-based loan pricing features are computed.")
    ]
    for t, d in points:
        p_t = tf_m.add_paragraph()
        p_t.text = f"• {t}: "
        p_t.font.bold = True
        p_t.font.size = Pt(12)
        p_t.font.color.rgb = COLOR_CYAN
        p_t.space_before = Pt(10)

        p_d = tf_m.add_paragraph()
        p_d.text = d
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = COLOR_TEXT_MUTED

    # --------------------------------------------------------------------------
    # SLIDE 11: PILLAR 5 — CONTINUOUS MONITORING & CYBER RESILIENCE
    # --------------------------------------------------------------------------
    slide11 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide11, COLOR_NAVY_DARK)
    add_header(slide11, "Pillar 5", "Continuous MLOps Monitoring & Cyber Resilience")

    monitoring_items = [
        ("📉 Data & Concept Drift Monitoring", "Detect when macroeconomic conditions alter member repayment behaviors, triggering model retraining alerts.", COLOR_GOLD),
        ("🛡️ Threat Modeling & Adversarial Defense", "Protect AI endpoints against data poisoning, prompt injection, and unauthorized parameter extraction.", COLOR_CYAN),
        ("📑 Cryptographic Event Audit Logging", "Store model input/output decision receipts in immutable, tamper-evident audit logs.", COLOR_EMERALD)
    ]

    for i, (m_t, m_d, m_c) in enumerate(monitoring_items):
        y = Inches(1.8 + i * 1.7)
        box = slide11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, Inches(11.733), Inches(1.4))
        box.fill.solid()
        box.fill.fore_color.rgb = COLOR_NAVY_CARD
        box.line.color.rgb = m_c

        tb_i = slide11.shapes.add_textbox(Inches(1.1), y + Inches(0.2), Inches(11.1), Inches(1.0))
        tf_i = tb_i.text_frame
        tf_i.word_wrap = True

        p = tf_i.paragraphs[0]
        p.text = m_t
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = m_c

        p_d = tf_i.add_paragraph()
        p_d.text = m_d
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = COLOR_TEXT_MUTED
        p_d.space_before = Pt(4)

    # --------------------------------------------------------------------------
    # SLIDE 12: BOARD OVERSIGHT & GOVERNANCE STRUCTURE
    # --------------------------------------------------------------------------
    slide12 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide12, COLOR_NAVY_DARK)
    add_header(slide12, "Board Oversight", "SACCO AI Governance Structure & Roles")

    roles = [
        ("👑 Board Audit & Risk Committee", "• Sets overall AI risk appetite\n• Reviews quarterly model audit reports\n• Approves high-risk AI deployments", COLOR_GOLD),
        ("🛡️ AI Governance Officer (AIECO)", "• Enforces KDPA 2019 compliance\n• Coordinates bias & fairness testing\n• Liaison for SASRA regulator audits", COLOR_CYAN),
        ("⚙️ IT & Credit Committee", "• Manages model deployment MLOps\n• Handles member appeal overrides\n• Monitors real-time model accuracy", COLOR_EMERALD)
    ]

    for i, (r_t, r_d, r_c) in enumerate(roles):
        x = Inches(0.8 + i * 4.0)
        box = slide12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.8), Inches(3.7), Inches(5.0))
        box.fill.solid()
        box.fill.fore_color.rgb = COLOR_NAVY_CARD
        box.line.color.rgb = r_c
        box.line.width = Pt(1.5)

        tb_r = slide12.shapes.add_textbox(x + Inches(0.2), Inches(2.0), Inches(3.3), Inches(4.5))
        tf_r = tb_r.text_frame
        tf_r.word_wrap = True

        p = tf_r.paragraphs[0]
        p.text = r_t
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = r_c

        p_d = tf_r.add_paragraph()
        p_d.text = r_d
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = COLOR_TEXT_MUTED
        p_d.space_before = Pt(10)

    # --------------------------------------------------------------------------
    # SLIDE 13: IMPLEMENTATION ROADMAP & STRATEGIC RECOMMENDATION
    # --------------------------------------------------------------------------
    slide13 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide13, COLOR_NAVY_DARK)
    add_header(slide13, "Implementation Roadmap", "Strategic Roadmap & Next Steps for the Board")

    roadmap = [
        ("Phase 1: Months 1 – 2", "Mobile App API Integration & Quick Wins", "Integrate STRIDE E-Voting tile & AI Credit Scoring engine into SACCO Phone App."),
        ("Phase 2: Months 3 – 4", "Governance Policy & SASRA Filing", "Formalize SACCO AI Charter & establish member consent flows."),
        ("Phase 3: Months 5 – 6", "XAI & Monitoring Dashboard", "Deploy Explainable AI dashboards & automated drift monitors.")
    ]

    for i, (ph_t, ph_n, ph_d) in enumerate(roadmap):
        y = Inches(1.8 + i * 1.5)
        box = slide13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, Inches(11.733), Inches(1.3))
        box.fill.solid()
        box.fill.fore_color.rgb = COLOR_NAVY_CARD
        box.line.color.rgb = COLOR_GOLD if i == 0 else (COLOR_CYAN if i == 1 else COLOR_EMERALD)

        tb_rd = slide13.shapes.add_textbox(Inches(1.1), y + Inches(0.15), Inches(11.1), Inches(1.0))
        tf_rd = tb_rd.text_frame
        tf_rd.word_wrap = True

        p = tf_rd.paragraphs[0]
        p.text = f"{ph_t} — {ph_n}"
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = COLOR_GOLD if i == 0 else (COLOR_CYAN if i == 1 else COLOR_EMERALD)

        p_d = tf_rd.add_paragraph()
        p_d.text = ph_d
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = COLOR_TEXT_MUTED
        p_d.space_before = Pt(4)

    # Bottom Recommendation Box
    rec_box = slide13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.1), Inches(11.733), Inches(0.9))
    rec_box.fill.solid()
    rec_box.fill.fore_color.rgb = COLOR_GOLD_DARK
    rec_box.line.color.rgb = COLOR_GOLD

    tb_rec = slide13.shapes.add_textbox(Inches(1.0), Inches(6.15), Inches(11.3), Inches(0.8))
    tf_rec = tb_rec.text_frame
    tf_rec.word_wrap = True

    p_rec = tf_rec.paragraphs[0]
    p_rec.text = "🎯 BOARD ACTION ITEM: Authorize integration of STRIDE E-Voting & AI Scoring API into Banki Kuu SACCO's existing Mobile Phone App."
    p_rec.font.size = Pt(13)
    p_rec.font.bold = True
    p_rec.font.color.rgb = COLOR_WHITE

    output_path = r"C:\Users\user\.gemini\antigravity\brain\13019df5-c73e-40a7-9dbf-48db7e5f7a5e\SACCO_AI_Governance_Mobile_App_Integration.pptx"
    prs.save(output_path)
    print(f"13-SLIDE DECK CREATED SUCCESSFULLY AT: {output_path}")

if __name__ == "__main__":
    create_deck()
