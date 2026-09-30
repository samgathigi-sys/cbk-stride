import os
import sys
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
from PIL import Image, ImageDraw, ImageFont, ImageFilter
import qrcode
from pptx import Presentation
from pptx.util import Inches
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
import shutil

DIR = os.path.dirname(os.path.abspath(__file__))
LOGO_PATH = os.path.join(DIR, "cbk_logo_transparent.png")
PORTAL_URL = "https://cbk-stride.streamlit.app/"

WIDTH = 1080
HEIGHT = 1920

# Executive Palette
C_NAVY_BG = (4, 15, 32)           # #040F20 Deep Midnight Sapphire Navy
C_CARD_BG = (9, 24, 48)           # #091830 Glassmorphic Sapphire Card
C_CARD_BORDER = (245, 197, 66)     # #F5C542 Championship Gold
C_CARD_BORDER_SUBTLE = (38, 72, 118) # Deep Slate Blue
C_GOLD = (245, 197, 66)           # #F5C542 Gold
C_GOLD_LIGHT = (255, 235, 143)    # Light Gold
C_CYAN = (0, 242, 254)            # #00F2FE Neon Telemetry Cyan
C_WHITE = (255, 255, 255)
C_SLATE = (148, 163, 184)         # Slate Gray
C_ICE = (226, 232, 240)           # High-legibility Ice White
C_EMERALD = (52, 211, 153)        # #34D399 Mint Green
C_AMBER = (251, 146, 60)          # #FB923C Amber Orange
C_PURPLE = (192, 132, 252)        # #C084FC Lavender

WIN_FONTS = "C:\\Windows\\Fonts"
def get_font(name, size):
    path = os.path.join(WIN_FONTS, name)
    if os.path.exists(path):
        try:
            return ImageFont.truetype(path, size)
        except Exception:
            pass
    return ImageFont.load_default()

font_title_hero = get_font("segoeuib.ttf", 52)
font_title = get_font("segoeuib.ttf", 36)
font_subtitle = get_font("segoeui.ttf", 20)
font_card_num = get_font("segoeuib.ttf", 16)
font_card_tag = get_font("segoeuib.ttf", 14)
font_card_h = get_font("segoeuib.ttf", 24)
font_card_body = get_font("segoeui.ttf", 18)
font_card_bullet = get_font("segoeuib.ttf", 18)
font_bottom_h = get_font("segoeuib.ttf", 20)
font_bottom_sub = get_font("segoeui.ttf", 15)
font_badge = get_font("segoeuib.ttf", 13)

print("Building Team Captain Field Roll-Call Quick Guide...")
flyer = Image.new("RGBA", (WIDTH, HEIGHT), C_NAVY_BG)
draw = ImageDraw.Draw(flyer)

# Ambient subtle gradient background
for y in range(450):
    intensity = 1.0 - (y / 450.0)
    alpha = int(35 * intensity)
    draw.line([(0, y), (WIDTH, y)], fill=(2, 91, 191, alpha))

for y in range(HEIGHT - 450, HEIGHT):
    intensity = (y - (HEIGHT - 450)) / 450.0
    alpha = int(30 * intensity)
    draw.line([(0, y), (WIDTH, y)], fill=(245, 197, 66, alpha))

# Framing Stripes
draw.rectangle([(0, 0), (WIDTH, 14)], fill=C_GOLD)
draw.rectangle([(0, HEIGHT - 14), (WIDTH, HEIGHT)], fill=C_GOLD)

# Top Header Section
top_y = 35

# Institutional Logo
if os.path.exists(LOGO_PATH):
    try:
        logo = Image.open(LOGO_PATH).convert("RGBA")
        logo.thumbnail((95, 95), Image.Resampling.LANCZOS)
        flyer.paste(logo, (55, top_y), logo)
    except Exception:
        pass

# Header Badges & Titles
draw.rounded_rectangle([(165, top_y + 4), (540, top_y + 28)], radius=5, fill=(9, 32, 64), outline=C_CARD_BORDER_SUBTLE, width=1)
draw.text((177, top_y + 8), "CENTRAL BANK OF KENYA  |  SPORTS SECRETARIAT", fill=C_GOLD, font=font_badge)

draw.text((165, top_y + 34), "CBK STRIDE", fill=C_WHITE, font=font_title_hero)
# TM superscipt-like pill
draw.rounded_rectangle([(448, top_y + 42), (488, top_y + 64)], radius=4, fill=C_GOLD)
draw.text((454, top_y + 44), "TM", fill=(4, 15, 32), font=get_font("segoeuib.ttf", 14))

draw.text((165, top_y + 102), "TEAM CAPTAIN'S ON-PITCH ROLL-CALL GUIDE", fill=C_CYAN, font=get_font("segoeuib.ttf", 22))
draw.text((165, top_y + 132), "Step-by-step field verification manual for appointed discipline captains.", fill=C_SLATE, font=get_font("segoeui.ttf", 16))

# Decorative line
draw.line([(55, top_y + 168), (WIDTH - 55, top_y + 168)], fill=C_CARD_BORDER_SUBTLE, width=1)

# ------------------------------------------------------------------------------
# PORTAL ACCESS HERO CARD WITH SCANNABLE QR CODE
# ------------------------------------------------------------------------------
qr_card_y = 225
qr_card_h = 240
draw.rounded_rectangle(
    [(55, qr_card_y), (WIDTH - 55, qr_card_y + qr_card_h)],
    radius=16,
    fill=(6, 20, 42),
    outline=C_CYAN,
    width=2
)

# Generate QR Code for Portal URL
qr = qrcode.QRCode(
    version=1,
    error_correction=qrcode.constants.ERROR_CORRECT_M,
    box_size=6,
    border=1
)
qr.add_data(PORTAL_URL)
qr.make(fit=True)
qr_img = qr.make_image(fill_color="#040F20", back_color="#FFFFFF").convert("RGBA")
qr_img = qr_img.resize((195, 195), Image.Resampling.LANCZOS)

# White rounded background for QR code
draw.rounded_rectangle([(75, qr_card_y + 22), (290, qr_card_y + 218)], radius=12, fill=C_WHITE)
flyer.paste(qr_img, (85, qr_card_y + 23), qr_img)

# QR Card Info Text
draw.rounded_rectangle([(315, qr_card_y + 25), (550, qr_card_y + 50)], radius=4, fill=(0, 242, 254, 40), outline=C_CYAN, width=1)
draw.text((327, qr_card_y + 29), "OFFICIAL CLOUD PORTAL", fill=C_CYAN, font=font_badge)

draw.text((315, qr_card_y + 58), "https://cbk-stride.streamlit.app/", fill=C_WHITE, font=get_font("segoeuib.ttf", 23))

draw.text((315, qr_card_y + 96), "1. Scan the QR code with any smartphone camera on the pitch.", fill=C_ICE, font=get_font("segoeui.ttf", 17))
draw.text((315, qr_card_y + 124), "2. Opens directly on Tab 1: [ Captain's Roll Call ] by default.", fill=C_ICE, font=get_font("segoeui.ttf", 17))
draw.text((315, qr_card_y + 152), "3. Follow the 3 fast steps below to clock in your squad in seconds.", fill=C_ICE, font=get_font("segoeui.ttf", 17))

draw.rounded_rectangle([(315, qr_card_y + 185), (630, qr_card_y + 212)], radius=5, fill=(16, 185, 129, 35), outline=C_EMERALD, width=1)
draw.text((327, qr_card_y + 189), "LIVE 24/7  -  MOBILE & TABLET READY", fill=C_EMERALD, font=font_badge)

# ------------------------------------------------------------------------------
# 3 STEP-BY-STEP FIELD ACTION CARDS
# ------------------------------------------------------------------------------
step_start_y = 495
step_card_h = 245
step_gap = 22

steps_data = [
    {
        "step": "STEP 01",
        "tag": "SECURITY & ISOLATION",
        "title": "Unlock Your Discipline Squad Sheet",
        "color": C_GOLD,
        "bullets": [
            ("Select Your Sport:", "Pick your discipline from the 18 Central Bank sports."),
            ("Enter Staff ID & Secret Passkey:", "Enter your Staff ID (e.g. 3428) & your private passkey."),
            ("Anti-Impersonation Active:", "Your squad unlocks and locks strictly to your sport. Cross-team tampering is 100% blocked by Central Bank security.")
        ]
    },
    {
        "step": "STEP 02",
        "tag": "KICKOFF VERIFICATION",
        "title": "Gate 1: Arrival Roll Call (On-Pitch Check-In)",
        "color": C_CYAN,
        "bullets": [
            ("Ensure Mode is 'Gate 1 (Arrival)':", "Select your station/venue post."),
            ("Method A (Rapid Batch Clock-In):", "Tap [ Select All Eligible ] then tap [ Batch Clock-In ]. All players are timestamped in 1 second!"),
            ("Method B (Individual Check):", "Tap [ Clock In Arrival ] next to each player standing before you on the grass/court.")
        ]
    },
    {
        "step": "STEP 03",
        "tag": "AUDIT CERTIFICATION",
        "title": "Gate 2: Departure (Clock-Out & Certification)",
        "color": C_EMERALD,
        "bullets": [
            ("Switch to 'Gate 2 (Departure)':", "When the match, drill, or practice concludes."),
            ("Clock Out Athletes:", "Tap departure for leaving players."),
            ("Automatic 45-Min Threshold:", "System pairs arrival and departure times, verifies minimum 45 mins active session, and awards certified attendance!")
        ]
    }
]

for idx, step in enumerate(steps_data):
    sy = step_start_y + idx * (step_card_h + step_gap)
    
    # Outer card
    draw.rounded_rectangle(
        [(55, sy), (WIDTH - 55, sy + step_card_h)],
        radius=14,
        fill=C_CARD_BG,
        outline=step["color"],
        width=2
    )

    # Top step pill
    draw.rounded_rectangle([(75, sy + 18), (175, sy + 44)], radius=5, fill=step["color"])
    draw.text((85, sy + 22), step["step"], fill=(4, 15, 32), font=font_card_num)

    # Category tag
    draw.rounded_rectangle([(185, sy + 18), (410, sy + 44)], radius=5, fill=(15, 34, 64), outline=C_CARD_BORDER_SUBTLE, width=1)
    draw.text((197, sy + 23), step["tag"], fill=C_ICE, font=font_card_tag)

    # Title
    draw.text((75, sy + 54), step["title"], fill=C_WHITE, font=font_card_h)

    # Divider line inside card
    draw.line([(75, sy + 90), (WIDTH - 75, sy + 90)], fill=C_CARD_BORDER_SUBTLE, width=1)

    # Bullets
    by = sy + 104
    for b_title, b_desc in step["bullets"]:
        draw.text((75, by), ">", fill=step["color"], font=font_card_bullet)
        draw.text((95, by), b_title, fill=step["color"], font=font_card_bullet)
        
        # Calculate offset
        b_box = font_card_bullet.getbbox(b_title)
        title_w = b_box[2] - b_box[0] + 8
        draw.text((95 + title_w, by), b_desc, fill=C_ICE, font=font_card_body)
        by += 38

# ------------------------------------------------------------------------------
# 3 POWER TOOL CARDS (PITCH-SIDE UTILITIES)
# ------------------------------------------------------------------------------
utils_y = 1320
util_w = (WIDTH - 110 - 30) // 3
util_h = 240

utils_data = [
    {
        "icon": "[ + ]",
        "title": "Pitch Walk-Ins",
        "tag": "UNSCHEDULED GUESTS",
        "color": C_AMBER,
        "desc": "A colleague arrived who wasn't on roster? Open the walk-in drawer, enter Payroll # and Name to add them instantly right on the pitch."
    },
    {
        "icon": "[ CSV ]",
        "title": "Certified Export",
        "tag": "INSTANT AUDIT REPORT",
        "color": C_PURPLE,
        "desc": "1-Tap download of the official roll-call CSV report, complete with exact second timestamps and captain verification signature tags."
    },
    {
        "icon": "[ RST ]",
        "title": "Safe Squad Reset",
        "tag": "CLEAN FIELD SLATE",
        "color": C_CYAN,
        "desc": "Need a clean slate for a new session? Reset attendance logs for just your sport without touching any of the other 17 disciplines."
    }
]

for idx, u in enumerate(utils_data):
    ux = 55 + idx * (util_w + 15)
    draw.rounded_rectangle([(ux, utils_y), (ux + util_w, utils_y + util_h)], radius=12, fill=C_CARD_BG, outline=C_CARD_BORDER_SUBTLE, width=1)
    
    # Left accent line
    draw.rectangle([(ux, utils_y + 15), (ux + 4, utils_y + util_h - 15)], fill=u["color"])
    
    draw.text((ux + 16, utils_y + 16), u["icon"], fill=u["color"], font=get_font("segoeuib.ttf", 16))
    draw.text((ux + 65, utils_y + 16), u["tag"], fill=C_SLATE, font=get_font("segoeui.ttf", 12))
    draw.text((ux + 16, utils_y + 44), u["title"], fill=C_WHITE, font=get_font("segoeuib.ttf", 20))
    draw.line([(ux + 16, utils_y + 74), (ux + util_w - 16, utils_y + 74)], fill=C_CARD_BORDER_SUBTLE, width=1)
    
    # Wrap description text
    words = u["desc"].split()
    lines = []
    cur_line = ""
    for w in words:
        test_line = f"{cur_line} {w}".strip()
        bbox = font_card_body.getbbox(test_line)
        if (bbox[2] - bbox[0]) < (util_w - 32):
            cur_line = test_line
        else:
            lines.append(cur_line)
            cur_line = w
    if cur_line:
        lines.append(cur_line)
    
    ly = utils_y + 88
    for line in lines[:5]:
        draw.text((ux + 16, ly), line, fill=C_ICE, font=get_font("segoeui.ttf", 15))
        ly += 25

# ------------------------------------------------------------------------------
# BOTTOM GOVERNANCE & AUDIT BAR
# ------------------------------------------------------------------------------
bottom_y = 1600
bottom_h = 160
draw.rounded_rectangle([(55, bottom_y), (WIDTH - 55, bottom_y + bottom_h)], radius=12, fill=(6, 20, 42), outline=C_CARD_BORDER_SUBTLE, width=1)

draw.text((75, bottom_y + 20), "CENTRAL BANK OF KENYA  |  ATHLETIC ACCREDITATION & GOVERNANCE", fill=C_GOLD, font=font_badge)
draw.text((75, bottom_y + 44), "Kenya Data Protection Act 2019 & Internal Audit Sovereign Standard", fill=C_WHITE, font=get_font("segoeuib.ttf", 20))
draw.text((75, bottom_y + 76), "All roll call check-ins are cryptographically hashed and permanently attributed to your captain credential.", fill=C_SLATE, font=get_font("segoeui.ttf", 15))
draw.text((75, bottom_y + 104), "Secretariat Support Hotline: Ext. 2408 / 2400  |  Email: sports.secretariat@centralbank.go.ke", fill=C_CYAN, font=get_font("segoeui.ttf", 15))

# Footer Stamp
draw.text((55, 1870), "CBK STRIDE (TM)  *  CENTRAL BANK OF KENYA SPORTS & WELLNESS SECRETARIAT  *  STRICTLY OFFICIAL", fill=C_SLATE, font=get_font("segoeui.ttf", 13))

# Save image assets
out_png = os.path.join(DIR, "CBK_STRIDE_Captain_RollCall_Guide.png")
out_jpg = os.path.join(DIR, "CBK_STRIDE_Captain_RollCall_Guide.jpg")
flyer.save(out_png, "PNG")
flyer.convert("RGB").save(out_jpg, "JPEG", quality=95)
print(f"Saved Image: {out_png}")
print(f"Saved Image: {out_jpg}")

# ------------------------------------------------------------------------------
# GENERATE PDF
# ------------------------------------------------------------------------------
out_pdf = os.path.join(DIR, "CBK_STRIDE_Captain_RollCall_Guide.pdf")
c = canvas.Canvas(out_pdf, pagesize=letter)
page_w, page_h = letter
c.drawImage(out_jpg, 0, 0, width=page_w, height=page_h)
c.showPage()
c.save()
print(f"Saved PDF: {out_pdf}")

# ------------------------------------------------------------------------------
# GENERATE PPTX
# ------------------------------------------------------------------------------
out_pptx = os.path.join(DIR, "CBK_STRIDE_Captain_RollCall_Guide.pptx")
prs = Presentation()
prs.slide_width = Inches(8.5)
prs.slide_height = Inches(15.11)
blank_slide_layout = prs.slide_layouts[6]
slide = prs.slides.add_slide(blank_slide_layout)
slide.shapes.add_picture(out_jpg, Inches(0), Inches(0), width=Inches(8.5), height=Inches(15.11))
prs.save(out_pptx)
print(f"Saved PPTX: {out_pptx}")

# Copy to Artifacts directory
ART_DIR = r"C:\Users\gathigisn.CBK.008\.gemini\antigravity\brain\00940864-5fa3-44bf-b1d9-02e121e7a052"
for f in [out_png, out_jpg, out_pdf, out_pptx]:
    try:
        shutil.copy(f, ART_DIR)
        print(f"Copied to Artifacts: {os.path.basename(f)}")
    except Exception as e:
        print(f"Artifact copy error: {e}")

print("Captain Roll Call Guide Flyer complete!")
