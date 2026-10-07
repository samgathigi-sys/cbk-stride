import os
import sys
import math
import json
import qrcode
import hmac
import hashlib
import openpyxl
from PIL import Image, ImageDraw, ImageFont, ImageFilter

BASE_DIR = r"C:\Users\user\.gemini\antigravity\scratch\cbk-stride"
EXCEL_PATH = r"C:\Users\user\.gemini\antigravity\brain\13019df5-c73e-40a7-9dbf-48db7e5f7a5e\.user_uploaded\media_1791364617995.xlsx"
TEMPLATE_FRONT = os.path.join(BASE_DIR, "TEMPLATE_MOMBASA_FRONT.png")
TEMPLATE_BACK = os.path.join(BASE_DIR, "TEMPLATE_MOMBASA_BACK.png")

OUTPUT_FRONT_DIR = os.path.join(BASE_DIR, "FRONT_CARDS")
OUTPUT_BACK_DIR = os.path.join(BASE_DIR, "BACK_CARDS")
OUTPUT_QR_DIR = os.path.join(BASE_DIR, "QR_CODES")
PREVIEW_DIR = r"C:\Users\user\.gemini\antigravity\brain\13019df5-c73e-40a7-9dbf-48db7e5f7a5e\badges_preview"
PRESS_DIR = os.path.join(BASE_DIR, "PRESS_READY_OUTPUT")

os.makedirs(OUTPUT_FRONT_DIR, exist_ok=True)
os.makedirs(OUTPUT_BACK_DIR, exist_ok=True)
os.makedirs(OUTPUT_QR_DIR, exist_ok=True)
os.makedirs(PREVIEW_DIR, exist_ok=True)
os.makedirs(PRESS_DIR, exist_ok=True)

HMAC_SECRET = b"BKS_MOMBASA_RETREAT_2026_ODPC_SECURE_TOKEN_SALT"

def get_font(size, bold=False):
    win_fonts = os.environ.get("WINDIR", r"C:\Windows") + r"\Fonts"
    name = "segoeuib.ttf" if bold else "segoeui.ttf"
    path = os.path.join(win_fonts, name)
    if os.path.exists(path):
        return ImageFont.truetype(path, size)
    return ImageFont.load_default()

def get_monogram(name):
    parts = [p for p in name.strip().split() if p]
    if len(parts) >= 2:
        return (parts[0][0] + parts[-1][0]).upper()
    elif len(parts) == 1:
        return parts[0][:2].upper()
    return "BK"

def render_mombasa_badge_pair(participant, front_tpl_img, back_tpl_img):
    front = front_tpl_img.copy()
    back = back_tpl_img.copy()
    cx = 489
    
    # ---------------- 1. FRONT BADGE ----------------
    draw_f = ImageDraw.Draw(front)
    
    # Monogram inside 3D Medallion
    mono = get_monogram(participant['name'])
    f_mono = get_font(132, bold=True)
    # Bevel shadow
    draw_f.text((cx + 4, 603), mono, fill=(4, 12, 28, 240), font=f_mono, anchor="mm")
    # Chrome face
    draw_f.text((cx, 598), mono, fill=(245, 250, 255, 255), font=f_mono, anchor="mm")
    # Top-left specular highlight
    draw_f.text((cx - 3, 595), mono, fill=(255, 255, 255, 200), font=f_mono, anchor="mm")
    
    # Participant Name
    name_str = participant['name'].strip().upper()
    if len(name_str) > 24:
        f_name = get_font(42, bold=True)
    elif len(name_str) > 18:
        f_name = get_font(46, bold=True)
    else:
        f_name = get_font(52, bold=True)
    draw_f.text((cx, 868), name_str, fill=(20, 48, 86), font=f_name, anchor="mm")
    
    # Designation Pill
    role_str = participant['role'].strip().upper()
    if len(role_str) > 28:
        f_pill = get_font(21, bold=True)
    else:
        f_pill = get_font(24, bold=True)
    bbox = f_pill.getbbox(role_str)
    pw = max(240, (bbox[2] - bbox[0]) + 54)
    ph = 46
    px0, py0 = cx - pw//2, 936 - ph//2
    draw_f.rounded_rectangle([(px0, py0), (px0 + pw, py0 + ph)], radius=23, fill=(24, 55, 96))
    draw_f.text((cx, 936), role_str, fill=(255, 255, 255), font=f_pill, anchor="mm")
    
    # Member No & Masked Staff ID
    f_meta = get_font(24, bold=False)
    sno_val = str(participant.get('sno', '1')).strip().zfill(3)
    mem_no = f"Member No: BKS-2026-{sno_val}"
    id_val = str(participant.get('id_no', '******02')).strip().replace('.0', '')
    id_masked = f"Staff ID: ******{id_val[-2:]}" if len(id_val)>=2 else "Staff ID: ******02"
    draw_f.text((cx, 1008), mem_no, fill=(65, 90, 120), font=f_meta, anchor="mm")
    draw_f.text((cx, 1046), id_masked, fill=(65, 90, 120), font=f_meta, anchor="mm")
    
    # Cryptographic Token & Front QR Code
    seed = f"BKS:{participant.get('sno')}:{participant['id_no']}:{participant['name']}".encode('utf-8')
    token_hash = hmac.new(HMAC_SECRET, seed, hashlib.sha256).hexdigest().upper()
    token_id = f"BKS-MSA26-{token_hash[:8]}"
    qr_payload = {
        "iss": "BANKI_KUU_SACCO",
        "evt": "MSA_RETREAT_2026",
        "tid": token_id,
        "sig": token_hash[:16],
        "sec": "ODPC_SEC_25",
        "v": 1
    }
    
    qr = qrcode.QRCode(
        version=1,
        error_correction=qrcode.constants.ERROR_CORRECT_H,
        box_size=6,
        border=1,
    )
    qr.add_data(json.dumps(qr_payload, separators=(',', ':')))
    qr.make(fit=True)
    qr_front_img = qr.make_image(fill_color=(18, 42, 75), back_color=(255, 255, 255)).resize((185, 185))
    
    # Pure White Card for Front QR
    qx0, qy0 = cx - 140, 1105
    qw, qh = 280, 240
    draw_f.rounded_rectangle([(qx0, qy0), (qx0 + qw, qy0 + qh)], radius=14, fill=(255, 255, 255), outline=(180, 205, 230), width=2)
    front.paste(qr_front_img, (cx - 185//2, qy0 + 10))
    f_qr_sub = get_font(18, bold=True)
    draw_f.text((cx, qy0 + qh - 18), "SCAN TO CONNECT", fill=(24, 55, 96), font=f_qr_sub, anchor="mm")
    
    # ---------------- 2. BACK BADGE ----------------
    draw_b = ImageDraw.Draw(back)
    bcx = 500
    
    f_b_dept = get_font(26, bold=True)
    f_b_name = get_font(34, bold=True)
    f_b_contact = get_font(24, bold=False)
    
    dept_label = participant.get('dept', 'Secretariat Section')
    draw_b.text((bcx, 635), dept_label, fill=(200, 225, 250), font=f_b_dept, anchor="mm")
    draw_b.text((bcx, 690), participant['name'], fill=(255, 255, 255), font=f_b_name, anchor="mm")
    
    phone = str(participant.get('phone_no', '+254 700 000 000')).strip().replace('.0', '')
    if phone.startswith("254") and len(phone) >= 12:
        phone_masked = f"+{phone[:3]} {phone[3:6]} *** *{phone[-2:]}"
    elif len(phone) >= 9:
        phone_masked = f"+254 {phone[-9:-6]} *** *{phone[-2:]}"
    else:
        phone_masked = "+254 7** *** ***"
    draw_b.text((bcx, 745), phone_masked, fill=(210, 230, 250), font=f_b_contact, anchor="mm")
    
    email = str(participant.get('email', '')).strip().rstrip(';')
    if "@" in email:
        user, domain = email.split("@", 1)
        masked_user = f"{user[0]}***{user[-1]}" if len(user) > 2 else f"{user[0]}***"
        email_clean = f"{masked_user}@{domain}"
    else:
        email_clean = "bks-tb@stride.co.ke"
    draw_b.text((bcx, 790), email_clean, fill=(180, 210, 240), font=f_b_contact, anchor="mm")
    
    # Back QR Code (x=195, y=965, 245x245)
    qr_back_img = qr.make_image(fill_color=(15, 35, 65), back_color=(255, 255, 255)).resize((245, 245))
    back.paste(qr_back_img, (195, 965))
    
    return front, back, qr_front_img, token_id

def load_all_95_participants(excel_path):
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    delegates = []
    
    # 1. BOD, Supervisory, Branch Reps (29)
    ws1 = wb['BOD,SC,BRANCH REPS & BHC BOD']
    for r in ws1.iter_rows(values_only=True):
        if r[0] is not None and str(r[0]).strip().isdigit():
            role_val = str(r[3]).strip() if r[3] and str(r[3]).strip() not in ('', 'None') else 'BRANCH REP'
            delegates.append({
                "sno": str(r[1]).strip(),
                "name": str(r[2]).strip(),
                "role": role_val,
                "id_no": str(r[4]).strip(),
                "phone_no": str(r[5]).strip(),
                "email": str(r[6] if len(r)>6 else '').strip(),
                "tier": "GOVERNANCE",
                "dept": "Governance & Leadership"
            })
            
    # 2. Secretariat Leadership (24)
    ws2 = wb['SECRETARIAT']
    for r in ws2.iter_rows(values_only=True):
        if r[0] is not None and str(r[0]).strip().isdigit():
            role_val = str(r[3]).strip() if r[3] and str(r[3]).strip() not in ('', 'None') else 'SECRETARIAT LEAD'
            delegates.append({
                "sno": str(r[1]).strip(),
                "name": str(r[2]).strip(),
                "role": role_val,
                "id_no": str(r[4]).strip(),
                "phone_no": str(r[5]).strip(),
                "email": str(r[6] if len(r)>6 else '').strip(),
                "tier": "SECRETARIAT",
                "dept": "Secretariat Section"
            })
            
    # 3. Operations, Canteen & Logistics (42)
    ws3 = wb['CANTEEN']
    for r in ws3.iter_rows(values_only=True):
        if r[0] is not None and str(r[0]).strip().isdigit():
            role_val = str(r[3]).strip() if r[3] and str(r[3]).strip() not in ('', 'None') else 'OPERATIONS & LOGISTICS'
            delegates.append({
                "sno": str(r[1]).strip(),
                "name": str(r[2]).strip(),
                "role": role_val,
                "id_no": str(r[4]).strip(),
                "phone_no": str(r[5]).strip(),
                "email": str(r[6] if len(r)>6 else '').strip(),
                "tier": "OPERATIONS",
                "dept": "Operations & Hospitality"
            })
            
    return delegates

def run():
    print("Loading clean Mombasa templates...")
    front_tpl = Image.open(TEMPLATE_FRONT).convert("RGBA")
    back_tpl = Image.open(TEMPLATE_BACK).convert("RGBA")
    
    print(f"Loading participants from {EXCEL_PATH}...")
    participants = load_all_95_participants(EXCEL_PATH)
    print(f"Loaded {len(participants)} delegates.")
    
    # First: Render Fred Opondo demo
    fred = {
        "sno": "031",
        "name": "Fred Opondo",
        "role": "HEAD IT",
        "id_no": "22345602",
        "phone_no": "254700000000",
        "email": "bks-tb@stride.co.ke",
        "dept": "Secretariat Section"
    }
    f_fred, b_fred, _, _ = render_mombasa_badge_pair(fred, front_tpl, back_tpl)
    f_fred.save(os.path.join(PREVIEW_DIR, "mombasa_fred_opondo_front.png"), "PNG", dpi=(300, 300))
    b_fred.save(os.path.join(PREVIEW_DIR, "mombasa_fred_opondo_back.png"), "PNG", dpi=(300, 300))
    print("Rendered Fred Opondo preview matching uploaded mockup!")
    
    # Second: Render CEO Albert Onchiri Orero
    ceo = participants[12]
    f_ceo, b_ceo, _, _ = render_mombasa_badge_pair(ceo, front_tpl, back_tpl)
    f_ceo.save(os.path.join(PREVIEW_DIR, "mombasa_ceo_front.png"), "PNG", dpi=(300, 300))
    b_ceo.save(os.path.join(PREVIEW_DIR, "mombasa_ceo_back.png"), "PNG", dpi=(300, 300))
    print(f"Rendered CEO preview: {ceo['name']}")
    
    # Third: Batch generate all 95 attendees
    all_fronts = []
    all_backs = []
    
    for idx, p in enumerate(participants, 1):
        clean_name = "".join([c if c.isalnum() or c in ("-", "_") else "_" for c in p["name"]]).strip("_")
        front_img, back_img, qr_img, token_id = render_mombasa_badge_pair(p, front_tpl, back_tpl)
        
        # Save individual PNGs
        f_path = os.path.join(OUTPUT_FRONT_DIR, f"{idx:02d}_{p['sno']}_{clean_name}_front.png")
        b_path = os.path.join(OUTPUT_BACK_DIR, f"{idx:02d}_{p['sno']}_{clean_name}_back.png")
        qr_path = os.path.join(OUTPUT_QR_DIR, f"{idx:02d}_{p['sno']}_{clean_name}_qr.png")
        
        front_img.save(f_path, "PNG", dpi=(300, 300))
        back_img.save(b_path, "PNG", dpi=(300, 300))
        qr_img.save(qr_path, "PNG", dpi=(300, 300))
        
        all_fronts.append(front_img.convert("RGB"))
        all_backs.append(back_img.convert("RGB"))
        
        if idx % 15 == 0 or idx == len(participants):
            print(f"  Compiled [{idx:02d}/{len(participants)}] {p['name']} ({p['role']})")
            
    # Compile 190-page press ready duplex PDF
    duplex_pdf_path = os.path.join(PRESS_DIR, "CBK_MOMBASA_2026_PRESS_READY_BADGES_DUPLEX.pdf")
    duplex_pages = []
    for f, b in zip(all_fronts, all_backs):
        duplex_pages.append(f)
        duplex_pages.append(b)
        
    duplex_pages[0].save(
        duplex_pdf_path,
        "PDF",
        resolution=300.0,
        save_all=True,
        append_images=duplex_pages[1:]
    )
    print(f"Compiled Master Duplex PDF: {duplex_pdf_path} ({len(duplex_pages)} pages)")
    
    # Compile 4-Up Imposition PDF
    print("\nCompiling 4-Up Commercial Imposition Sheets (48 Pages at 300 DPI)...")
    SHEET_W, SHEET_H = 2480, 3508
    BW, BH = 1000, 1500
    left_m, center_gx = 160, 160
    top_m, center_gy = 200, 100
    
    positions = [
        (left_m, top_m),
        (left_m + BW + center_gx, top_m),
        (left_m, top_m + BH + center_gy),
        (left_m + BW + center_gx, top_m + BH + center_gy),
    ]
    
    def draw_crop_marks(draw, x0, y0, x1, y1, mark_len=26, offset=10, w=2):
        col = (110, 110, 110)
        draw.line([(x0 - offset - mark_len, y0), (x0 - offset, y0)], fill=col, width=w)
        draw.line([(x0, y0 - offset - mark_len), (x0, y0 - offset)], fill=col, width=w)
        draw.line([(x1 + offset, y0), (x1 + offset + mark_len, y0)], fill=col, width=w)
        draw.line([(x1, y0 - offset - mark_len), (x1, y0 - offset)], fill=col, width=w)
        draw.line([(x0 - offset - mark_len, y1), (x0 - offset, y1)], fill=col, width=w)
        draw.line([(x0, y1 + offset), (x0, y1 + offset + mark_len)], fill=col, width=w)
        draw.line([(x1 + offset, y1), (x1 + offset + mark_len, y1)], fill=col, width=w)
        draw.line([(x1, y1 + offset), (x1, y1 + offset + mark_len)], fill=col, width=w)
        
    def build_sheet(badges, sheet_num, total_s, is_back=False):
        sheet = Image.new("RGB", (SHEET_W, SHEET_H), (255, 255, 255))
        sdraw = ImageDraw.Draw(sheet)
        order = [1, 0, 3, 2] if is_back else [0, 1, 2, 3]
        for slot, b_idx in enumerate(order):
            if b_idx < len(badges) and badges[b_idx] is not None:
                bx, by = positions[slot]
                b_resized = badges[b_idx].resize((BW, BH), Image.Resampling.LANCZOS)
                sheet.paste(b_resized, (bx, by))
                draw_crop_marks(sdraw, bx, by, bx + BW, by + BH)
        # Header slug
        side = "BACK SHEET (DUPLEX FLIPPED)" if is_back else "FRONT SHEET"
        f_slug = get_font(18, bold=True)
        slug = f"BANKI KUU SACCO RETREAT 2026 • MOMBASA • PRESS SHEET {sheet_num:02d}/{total_s:02d} [{side}] • 300 DPI"
        sdraw.text((SHEET_W // 2, 130), slug, fill=(90, 90, 90), font=f_slug, anchor="mm")
        return sheet
        
    imposition_sheets = []
    total_sheets = math.ceil(len(participants) / 4)
    for s_idx in range(total_sheets):
        start_p = s_idx * 4
        c_front = all_fronts[start_p:start_p + 4]
        c_back = all_backs[start_p:start_p + 4]
        while len(c_front) < 4:
            c_front.append(None)
            c_back.append(None)
            
        fs = build_sheet(c_front, s_idx + 1, total_sheets, is_back=False)
        bs = build_sheet(c_back, s_idx + 1, total_sheets, is_back=True)
        imposition_sheets.append(fs)
        imposition_sheets.append(bs)
        
        if s_idx == 0:
            fs.save(os.path.join(PREVIEW_DIR, "sheet_01_front_imposition_preview.png"), "PNG", dpi=(300, 300))
            bs.save(os.path.join(PREVIEW_DIR, "sheet_01_back_imposition_preview.png"), "PNG", dpi=(300, 300))
            
    imposition_pdf_path = os.path.join(PRESS_DIR, "CBK_MOMBASA_2026_IMPOSITION_SHEETS_4UP.pdf")
    imposition_sheets[0].save(
        imposition_pdf_path,
        "PDF",
        resolution=300.0,
        save_all=True,
        append_images=imposition_sheets[1:]
    )
    print(f"Compiled Master 4-Up Imposition PDF: {imposition_pdf_path} ({len(imposition_sheets)} pages)")
    print("ALL 95 MOMBASA BADGES & IMPOSITION SHEETS GENERATED SUCCESSFULLY!")

if __name__ == "__main__":
    run()
